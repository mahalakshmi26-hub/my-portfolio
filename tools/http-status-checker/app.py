"""
HTTP Status Code Checker — Flask dashboard
Built by Mahalakshmi Marimuthu · Digital Marketing Strategist & AI-Powered SEO Expert

Takes a pasted list of URLs, checks the HTTP status code of each one (following
redirects, so the final destination and the full redirect chain are captured),
and reports a health score plus a downloadable CSV / Excel report.
"""

import io
import random
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urlparse

import requests
from flask import Flask, jsonify, render_template_string, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# ---------------------------------------------------------------------------
# Config / limits (keep checks fast + safe on Render's free tier)
# ---------------------------------------------------------------------------
MAX_URLS = 200                  # how many URLs we'll check in a single run
CHECK_TIMEOUT = 8
WORKERS = 20
DOMAIN_CONCURRENCY = 6          # max simultaneous requests to any single domain
                                 # (avoids tripping bot/WAF protection with a burst)
# A normal browser UA — self-identifying as a bot gets flagged/blocked by most
# WAFs (Akamai, Cloudflare, etc.) before a single request even completes.
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)
HEADERS = {
    "User-Agent": USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}

_domain_semaphores = {}
_domain_semaphores_lock = threading.Lock()


def _get_domain_semaphore(url):
    """One semaphore per host, so we never hammer a single domain with a burst
    of concurrent requests (a common trigger for WAF/bot-protection 403s)."""
    netloc = urlparse(url).netloc.lower()
    with _domain_semaphores_lock:
        sem = _domain_semaphores.get(netloc)
        if sem is None:
            sem = threading.Semaphore(DOMAIN_CONCURRENCY)
            _domain_semaphores[netloc] = sem
        return sem


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def normalize_url(url):
    """Strip whitespace + fragment (fragments never reach the server anyway).
    Deliberately does NOT touch a trailing slash — that can change the
    response a server gives, so we check exactly what was pasted."""
    url = url.strip()
    if not url:
        return ""
    return url.split("#")[0].strip()


def is_checkable_scheme(url):
    scheme = urlparse(url).scheme.lower()
    return scheme in ("http", "https")


# Domains known to routinely reject simple HTTP requests (no cookies/JS/browser
# fingerprint) even when the page is genuinely live — social platforms and URL
# shorteners are the classic offenders. A failed check against one of these is
# reported as "unverifiable" instead of a real 4xx/5xx/error.
UNRELIABLE_DOMAINS = {
    "facebook.com", "instagram.com", "twitter.com", "x.com", "linkedin.com",
    "threads.net", "tiktok.com",
    "bit.ly", "tinyurl.com", "t.co", "goo.gl", "rebrand.ly", "ow.ly", "bnkbzr.co",
}


def is_unreliable_domain(url):
    netloc = urlparse(url).netloc.lower()
    netloc = netloc[4:] if netloc.startswith("www.") else netloc
    return any(netloc == d or netloc.endswith("." + d) for d in UNRELIABLE_DOMAINS)


def classify(status_code, redirect_count, error):
    if error or status_code is None:
        return "error"
    if 200 <= status_code < 300:
        return "redirect" if redirect_count > 0 else "success"
    if 300 <= status_code < 400:
        return "redirect"
    if 400 <= status_code < 500:
        return "client_error"
    if 500 <= status_code < 600:
        return "server_error"
    return "error"


def _request_once(url):
    resp = requests.head(url, headers=HEADERS, timeout=CHECK_TIMEOUT, allow_redirects=True)
    # A lot of WAFs (Akamai, Cloudflare, etc.) treat a bare HEAD request as a bot
    # signal and block it specifically, while a normal GET sails through. So any
    # 4xx/5xx from HEAD gets a GET before we trust it — cheap insurance, and a
    # genuinely dead page will still come back the same way on GET.
    if resp.status_code >= 400:
        resp = requests.get(
            url, headers=HEADERS, timeout=CHECK_TIMEOUT, allow_redirects=True, stream=True
        )
        resp.close()
    return resp


def check_url(url):
    start = time.time()
    sem = _get_domain_semaphore(url)
    sem.acquire()
    try:
        resp = _request_once(url)
        # 403/429/503 are frequently a bot-protection layer blocking a burst of
        # requests rather than a genuinely broken URL — pause briefly and
        # retry once before trusting it.
        if resp.status_code in (403, 429, 503):
            time.sleep(1.2 + random.random())
            resp = _request_once(url)
        elapsed_ms = round((time.time() - start) * 1000)
        redirect_chain = [{"status_code": h.status_code, "url": h.url} for h in resp.history]
        severity = classify(resp.status_code, len(resp.history), None)
        if severity in ("client_error", "server_error") and is_unreliable_domain(url):
            severity = "unverifiable"
        return {
            "status_code": resp.status_code,
            "final_url": resp.url,
            "redirect_count": len(resp.history),
            "redirect_chain": redirect_chain,
            "response_ms": elapsed_ms,
            "error": None,
            "severity": severity,
        }
    except requests.exceptions.Timeout:
        return _error_result(url, start, "Timeout")
    except requests.exceptions.TooManyRedirects:
        return _error_result(url, start, "Too many redirects")
    except requests.exceptions.SSLError:
        return _error_result(url, start, "SSL error")
    except requests.exceptions.ConnectionError:
        return _error_result(url, start, "Connection failed")
    except Exception as exc:
        return _error_result(url, start, str(exc)[:120])
    finally:
        sem.release()


def _error_result(url, start, message):
    severity = "unverifiable" if is_unreliable_domain(url) else "error"
    return {
        "status_code": None,
        "final_url": None,
        "redirect_count": 0,
        "redirect_chain": [],
        "response_ms": round((time.time() - start) * 1000),
        "error": message,
        "severity": severity,
    }


def run_check(urls):
    results = []
    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures = {pool.submit(check_url, u): u for u in urls}
        for fut in as_completed(futures):
            url = futures[fut]
            results.append({"url": url, **fut.result()})

    # Worst-first ordering so problems surface at the top of the table.
    severity_rank = {"error": 0, "server_error": 1, "client_error": 2, "unverifiable": 3, "redirect": 4, "success": 5}
    results.sort(key=lambda r: severity_rank.get(r["severity"], 6))

    total = len(results)
    success = sum(1 for r in results if r["severity"] == "success")
    redirect = sum(1 for r in results if r["severity"] == "redirect")
    client_error = sum(1 for r in results if r["severity"] == "client_error")
    server_error = sum(1 for r in results if r["severity"] == "server_error")
    error = sum(1 for r in results if r["severity"] == "error")
    unverifiable = sum(1 for r in results if r["severity"] == "unverifiable")
    scoreable = total - unverifiable
    health_score = round((success / scoreable) * 100) if scoreable else 100

    status_counts = {}
    for r in results:
        code = r["status_code"]
        key = str(code) if code is not None else (r["error"] or "Error")
        status_counts[key] = status_counts.get(key, 0) + 1
    top_status_codes = sorted(status_counts.items(), key=lambda kv: kv[1], reverse=True)[:8]

    summary = {
        "total": total,
        "success": success,
        "redirect": redirect,
        "client_error": client_error,
        "server_error": server_error,
        "error": error,
        "unverifiable": unverifiable,
        "health_score": health_score,
        "top_status_codes": top_status_codes,
    }
    return summary, results


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template_string(PAGE_TEMPLATE, max_urls=MAX_URLS)


@app.route("/api/check", methods=["POST"])
def api_check():
    data = request.get_json(force=True, silent=True) or {}
    raw_input = (data.get("input") or "").strip()

    if not raw_input:
        return jsonify({"error": "Please paste at least one URL."}), 400

    urls, seen = [], set()
    for line in re.split(r"[\n,]+", raw_input):
        u = normalize_url(line)
        if not u:
            continue
        if not is_checkable_scheme(u):
            u = "https://" + u
        if not is_checkable_scheme(u):
            continue
        if u not in seen:
            seen.add(u)
            urls.append(u)

    if not urls:
        return jsonify({"error": "No valid http(s) URLs found in what you pasted."}), 400

    truncated = len(urls) > MAX_URLS
    urls = urls[:MAX_URLS]

    try:
        summary, results = run_check(urls)
    except Exception as exc:
        return jsonify({"error": f"Check failed: {exc}"}), 500

    summary["truncated"] = truncated
    return jsonify({"summary": summary, "results": results})


@app.route("/api/export/excel", methods=["POST"])
def api_export_excel():
    data = request.get_json(force=True, silent=True) or {}
    results = data.get("results") or []
    if not results:
        return jsonify({"error": "No results to export."}), 400

    wb = Workbook()
    ws = wb.active
    ws.title = "HTTP Status Report"

    headers = ["URL", "Status Code", "Category", "Final URL", "Redirect Count", "Redirect Chain", "Error"]
    ws.append(headers)
    header_fill = PatternFill(start_color="7C6AF7", end_color="7C6AF7", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    category_labels = {
        "success": "Success (2xx)",
        "redirect": "Redirect (3xx)",
        "client_error": "Client Error (4xx)",
        "server_error": "Server Error (5xx)",
        "error": "Unreachable",
        "unverifiable": "Unverifiable",
    }
    category_fills = {
        "success": "E4FBF2",
        "redirect": "FFF1E4",
        "client_error": "FDE9E9",
        "server_error": "FCE0E0",
        "error": "F3E4E4",
        "unverifiable": "F0F0F5",
    }

    for r in results:
        chain = " -> ".join(
            f"{hop.get('status_code', '')} {hop.get('url', '')}".strip()
            for hop in (r.get("redirect_chain") or [])
        )
        severity = r.get("severity", "")
        row = [
            r.get("url", ""),
            r.get("status_code") if r.get("status_code") is not None else "",
            category_labels.get(severity, severity),
            r.get("final_url") or "",
            r.get("redirect_count", 0),
            chain,
            r.get("error") or "",
        ]
        ws.append(row)
        fill_color = category_fills.get(severity)
        if fill_color:
            row_idx = ws.max_row
            fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    widths = [55, 12, 18, 55, 14, 60, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="http-status-report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/healthz")
def healthz():
    return jsonify({"status": "ok"})


# ---------------------------------------------------------------------------
# Frontend (single-file dashboard template)
# ---------------------------------------------------------------------------
PAGE_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>HTTP Status Code Checker | SEO Tools by Mahalakshmi</title>
<meta name="description" content="Free HTTP status code checker. Paste a batch of URLs to check status codes, final destinations, and redirect chains, then export the report as CSV or Excel.">
<link rel="icon" type="image/svg+xml" href="data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 32 32'%3E%3Crect width='32' height='32' rx='8' fill='%237c6af7'/%3E%3Ctext x='50%25' y='50%25' dominant-baseline='central' text-anchor='middle' font-family='Georgia%2C serif' font-size='18' font-weight='900' fill='white'%3EM%3C/text%3E%3C/svg%3E">
<link href="https://fonts.googleapis.com/css2?family=Sora:wght@400;600;700;800;900&family=Inter:wght@300;400;500;600&family=DM+Mono:wght@400;500&display=swap" rel="stylesheet">
<script src="https://cdn.jsdelivr.net/npm/chart.js@4"></script>
<style>
  :root {
    --bg: #0a0a0f; --surface: #12121a; --surface2: #1a1a26;
    --accent: #7c6af7; --accent2: #f7a26a; --accent3: #6af7c8;
    --danger: #f76a6a; --text: #e8e8f0; --muted: #8888a8;
    --border: rgba(124,106,247,0.18); --card: rgba(18,18,28,0.85);
  }
  *, *::before, *::after { box-sizing: border-box; margin: 0; padding: 0; }
  body { font-family: 'Inter', sans-serif; background: var(--bg); color: var(--text); font-size: 15px; line-height: 1.6; }
  a { color: inherit; }

  .layout { display: flex; min-height: 100vh; }

  /* SIDEBAR */
  .sidebar { width: 240px; flex-shrink: 0; background: var(--surface); border-right: 1px solid var(--border); padding: 1.6rem 1.2rem; position: sticky; top: 0; height: 100vh; display: flex; flex-direction: column; gap: 2rem; }
  .brand { display: flex; align-items: center; gap: 0.55rem; text-decoration: none; }
  .brand-mark { width: 30px; height: 30px; border-radius: 8px; background: var(--accent); display: flex; align-items: center; justify-content: center; font-family: 'Sora', sans-serif; font-weight: 900; font-size: 1rem; color: #fff; flex-shrink: 0; }
  .brand-text { font-family: 'Sora', sans-serif; font-weight: 700; font-size: 0.95rem; }
  .sidebar-section { display: flex; flex-direction: column; gap: 0.3rem; }
  .sidebar-label { font-family: 'DM Mono', monospace; font-size: 0.65rem; letter-spacing: 0.14em; text-transform: uppercase; color: var(--muted); margin-bottom: 0.5rem; }
  .sidebar-link { display: flex; align-items: center; gap: 0.6rem; padding: 0.55rem 0.7rem; border-radius: 8px; border: 1px solid transparent; font-size: 0.85rem; font-weight: 500; color: var(--muted); text-decoration: none; transition: all 0.15s; }
  .sidebar-link:hover { background: rgba(124,106,247,0.12); color: var(--text); }
  .sidebar-link.active { background: rgba(124,106,247,0.14); border: 1px solid rgba(124,106,247,0.4); color: var(--accent); font-weight: 700; }
  .sidebar-limits { margin-top: auto; background: var(--card); border: 1px solid var(--border); border-radius: 10px; padding: 0.9rem; font-size: 0.72rem; color: var(--muted); line-height: 1.6; }
  .sidebar-limits strong { color: var(--text); }
  .sidebar-footer { padding-top: 1rem; border-top: 1px solid var(--border); }
  .sidebar-footer p { font-size: 0.7rem; color: var(--muted); line-height: 1.5; margin-bottom: 0.7rem; }
  .credit-name { color: var(--accent3); font-weight: 700; text-decoration: none; }
  .credit-name:hover { text-decoration: underline; }
  .sidebar-social { display: flex; gap: 0.5rem; }
  .sidebar-social a { width: 28px; height: 28px; border-radius: 6px; background: rgba(106,247,200,0.08); border: 1px solid rgba(106,247,200,0.35); display: flex; align-items: center; justify-content: center; color: var(--accent3); text-decoration: none; font-size: 0.68rem; font-family: 'DM Mono', monospace; transition: all 0.2s; }
  .sidebar-social a:hover { background: rgba(106,247,200,0.2); border-color: var(--accent3); }

  /* MAIN */
  .main { flex: 1; padding: 2.2rem 3rem 4rem; max-width: 1180px; }
  .page-head { margin-bottom: 1.8rem; }
  .page-eyebrow { font-family: 'DM Mono', monospace; font-size: 0.68rem; letter-spacing: 0.16em; text-transform: uppercase; color: var(--accent3); margin-bottom: 0.5rem; }
  .page-title { font-family: 'Sora', sans-serif; font-size: 1.7rem; font-weight: 800; margin-bottom: 0.4rem; }
  .page-sub { color: var(--muted); font-size: 0.92rem; max-width: 640px; }

  /* CHECK CARD */
  .scan-card { background: var(--card); border: 1px solid var(--border); border-radius: 14px; padding: 1.6rem; margin-bottom: 2rem; }
  .scan-input { width: 100%; background: var(--surface2); border: 1px solid var(--border); border-radius: 8px; color: var(--text); padding: 0.75rem 0.9rem; font-family: 'Inter', sans-serif; font-size: 0.88rem; resize: vertical; min-height: 150px; }
  .scan-hint { font-size: 0.74rem; color: var(--muted); margin-top: 0.5rem; font-family: 'DM Mono', monospace; }
  .scan-actions { margin-top: 1.2rem; display: flex; align-items: center; gap: 1rem; }
  .btn-primary { background: var(--accent); color: #fff; border: none; padding: 0.7rem 1.7rem; border-radius: 8px; font-family: 'Inter', sans-serif; font-weight: 600; font-size: 0.88rem; cursor: pointer; transition: background 0.2s, transform 0.2s; }
  .btn-primary:hover { background: #6a58e8; transform: translateY(-1px); }
  .btn-primary:disabled { opacity: 0.55; cursor: not-allowed; transform: none; }
  .spinner { width: 16px; height: 16px; border: 2px solid rgba(255,255,255,0.35); border-top-color: #fff; border-radius: 50%; animation: spin 0.7s linear infinite; display: inline-block; vertical-align: -3px; margin-right: 0.5rem; }
  @keyframes spin { to { transform: rotate(360deg); } }
  .scan-status { font-size: 0.82rem; color: var(--muted); }
  .error-box { margin-top: 1rem; background: rgba(247,106,106,0.1); border: 1px solid rgba(247,106,106,0.3); color: #ff9d9d; padding: 0.7rem 1rem; border-radius: 8px; font-size: 0.85rem; display: none; }

  /* RESULTS */
  #results { display: none; }
  .summary-row { display: grid; grid-template-columns: 200px repeat(5, 1fr); gap: 1rem; margin-bottom: 1.6rem; }
  .health-card { background: var(--card); border: 1px solid var(--border); border-radius: 14px; padding: 1.2rem; display: flex; flex-direction: column; align-items: center; justify-content: center; gap: 0.4rem; }
  .health-ring-wrap { position: relative; width: 100px; height: 100px; }
  .health-ring-wrap svg { transform: rotate(-90deg); }
  .health-ring-bg { fill: none; stroke: var(--surface2); stroke-width: 9; }
  .health-ring-fg { fill: none; stroke: var(--accent3); stroke-width: 9; stroke-linecap: round; transition: stroke-dashoffset 0.6s ease; }
  .health-score { position: absolute; inset: 0; display: flex; align-items: center; justify-content: center; font-family: 'Sora', sans-serif; font-size: 1.3rem; font-weight: 800; }
  .health-label { font-size: 0.7rem; color: var(--muted); font-family: 'DM Mono', monospace; text-transform: uppercase; letter-spacing: 0.08em; }
  .stat-card { background: var(--card); border: 1px solid var(--border); border-radius: 14px; padding: 1.1rem 1.2rem; display: flex; flex-direction: column; justify-content: center; gap: 0.3rem; }
  .stat-num { font-family: 'Sora', sans-serif; font-size: 1.6rem; font-weight: 800; }
  .stat-label { font-size: 0.75rem; color: var(--muted); }
  .stat-card.success .stat-num { color: var(--accent3); }
  .stat-card.redirect .stat-num { color: var(--accent2); }
  .stat-card.errors .stat-num { color: var(--danger); }
  .stat-card.unverifiable .stat-num { color: var(--muted); }

  .chart-row { display: grid; grid-template-columns: 1fr 1fr; gap: 1rem; margin-bottom: 1.6rem; }
  .chart-card { background: var(--card); border: 1px solid var(--border); border-radius: 14px; padding: 1.2rem; }
  .chart-card h3 { font-family: 'Sora', sans-serif; font-size: 0.85rem; font-weight: 700; margin-bottom: 0.9rem; }
  .chart-card canvas { max-height: 190px; }

  .results-card { background: var(--card); border: 1px solid var(--border); border-radius: 14px; padding: 1.2rem; }
  .results-head { display: flex; align-items: center; justify-content: space-between; flex-wrap: wrap; gap: 0.8rem; margin-bottom: 1rem; }
  .filter-row { display: flex; gap: 0.5rem; flex-wrap: wrap; }
  .filter-btn { background: var(--surface2); border: 1px solid var(--border); color: var(--muted); padding: 0.4rem 0.85rem; border-radius: 100px; font-size: 0.76rem; font-family: 'DM Mono', monospace; cursor: pointer; transition: all 0.15s; }
  .filter-btn.active { background: var(--accent); border-color: var(--accent); color: #fff; }
  .export-row { display: flex; gap: 0.5rem; }
  .btn-export { background: var(--surface2); border: 1px solid var(--border); color: var(--text); padding: 0.45rem 1rem; border-radius: 8px; font-size: 0.78rem; font-weight: 600; cursor: pointer; }
  .btn-export:hover { border-color: var(--accent); color: var(--accent); }
  .btn-export:disabled { opacity: 0.5; cursor: not-allowed; }

  table { width: 100%; border-collapse: collapse; font-size: 0.82rem; }
  thead th { text-align: left; padding: 0.6rem 0.7rem; color: var(--muted); font-size: 0.68rem; text-transform: uppercase; letter-spacing: 0.06em; font-family: 'DM Mono', monospace; border-bottom: 1px solid var(--border); }
  tbody td { padding: 0.65rem 0.7rem; border-bottom: 1px solid rgba(255,255,255,0.04); vertical-align: top; }
  tbody tr:hover { background: rgba(124,106,247,0.05); }
  .url-cell { max-width: 340px; overflow-wrap: anywhere; }
  .url-cell a { text-decoration: none; color: var(--text); }
  .url-cell a:hover { color: var(--accent); }
  .redirect-chain { font-size: 0.72rem; color: var(--muted); margin-top: 0.2rem; overflow-wrap: anywhere; }
  .badge { display: inline-block; font-family: 'DM Mono', monospace; font-size: 0.68rem; padding: 0.22rem 0.6rem; border-radius: 100px; font-weight: 600; white-space: nowrap; }
  .badge.success { background: rgba(106,247,200,0.12); color: var(--accent3); border: 1px solid rgba(106,247,200,0.25); }
  .badge.redirect { background: rgba(247,162,106,0.12); color: var(--accent2); border: 1px solid rgba(247,162,106,0.25); }
  .badge.client_error, .badge.server_error, .badge.error { background: rgba(247,106,106,0.12); color: var(--danger); border: 1px solid rgba(247,106,106,0.3); }
  .badge.unverifiable { background: rgba(136,136,168,0.15); color: var(--muted); border: 1px solid rgba(136,136,168,0.3); }
  .empty-state { text-align: center; padding: 3rem 1rem; color: var(--muted); font-size: 0.88rem; }

  @media (max-width: 900px) {
    .sidebar { display: none; }
    .main { padding: 1.6rem 1.2rem 3rem; }
    .summary-row { grid-template-columns: 1fr 1fr; }
    .chart-row { grid-template-columns: 1fr; }
  }
</style>
</head>
<body>
<div class="layout">

  <aside class="sidebar">
    <a href="https://mahalakshmi26-hub.github.io/my-portfolio/tools/index.html" class="brand">
      <span class="brand-mark">M</span>
      <span class="brand-text">SEO Tools</span>
    </a>
    <div class="sidebar-section">
      <div class="sidebar-label">Tool</div>
      <a href="#" class="sidebar-link active">🚦 HTTP Status Checker</a>
    </div>
    <div class="sidebar-section">
      <div class="sidebar-label">Navigate</div>
      <a href="https://mahalakshmi26-hub.github.io/my-portfolio/tools/index.html" class="sidebar-link">🧰 All My Tools</a>
      <a href="https://mahalakshmi26-hub.github.io/my-portfolio" class="sidebar-link">👩‍💻 Portfolio</a>
      <a href="https://mahalakshmi26-hub.github.io/my-portfolio/blog/index.html" class="sidebar-link">📝 Blog Posts</a>
    </div>
    <div class="sidebar-limits">
      Checks up to <strong>{{ max_urls }} URLs</strong> per run, so results stay fast and free.
    </div>
    <div class="sidebar-footer">
      <p>Built by <a href="https://mahalakshmi26-hub.github.io/my-portfolio" class="credit-name">Mahalakshmi Marimuthu</a><br>Digital Marketing Strategist &amp; AI-Powered SEO Expert</p>
      <div class="sidebar-social">
        <a href="https://linkedin.com/in/mahalakshmimarimuthu" target="_blank" title="LinkedIn">in</a>
        <a href="https://www.instagram.com/mahapravin26/" target="_blank" title="Instagram">IG</a>
        <a href="mailto:mahalakshmi.digitalpro@gmail.com" title="Email">✉</a>
      </div>
    </div>
  </aside>

  <main class="main">
    <div class="page-head">
      <div class="page-eyebrow">// seo tool</div>
      <div class="page-title">HTTP Status Code Checker</div>
      <div class="page-sub">Paste a batch of URLs to check their HTTP status codes, final destinations, and redirect chains in one go — with a health score and a downloadable report.</div>
    </div>

    <div class="scan-card">
      <textarea id="urlInput" class="scan-input" placeholder="https://example.com/page-1&#10;https://example.com/page-2&#10;https://example.com/page-3"></textarea>
      <div class="scan-hint">Paste one URL per line (or comma-separated) — up to {{ max_urls }} at a time.</div>

      <div class="scan-actions">
        <button class="btn-primary" id="checkBtn" onclick="startCheck()">Check Status Codes</button>
        <span class="scan-status" id="scanStatus"></span>
      </div>
      <div class="error-box" id="errorBox"></div>
    </div>

    <div id="results">
      <div class="summary-row">
        <div class="health-card">
          <div class="health-ring-wrap">
            <svg width="100" height="100" viewBox="0 0 100 100">
              <circle class="health-ring-bg" cx="50" cy="50" r="42"></circle>
              <circle class="health-ring-fg" id="healthRing" cx="50" cy="50" r="42" stroke-dasharray="264" stroke-dashoffset="264"></circle>
            </svg>
            <div class="health-score" id="healthScoreText">0%</div>
          </div>
          <div class="health-label">Status Health</div>
        </div>
        <div class="stat-card"><div class="stat-num" id="statTotal">0</div><div class="stat-label">URLs checked</div></div>
        <div class="stat-card success"><div class="stat-num" id="statSuccess">0</div><div class="stat-label">Success (2xx)</div></div>
        <div class="stat-card redirect"><div class="stat-num" id="statRedirect">0</div><div class="stat-label">Redirects (3xx)</div></div>
        <div class="stat-card errors"><div class="stat-num" id="statErrors">0</div><div class="stat-label">Errors (4xx/5xx)</div></div>
        <div class="stat-card unverifiable"><div class="stat-num" id="statUnverifiable">0</div><div class="stat-label">Unverifiable</div></div>
      </div>

      <div class="chart-row">
        <div class="chart-card">
          <h3>Status Category Breakdown</h3>
          <canvas id="statusChart"></canvas>
        </div>
        <div class="chart-card">
          <h3>Top Status Codes</h3>
          <canvas id="codesChart"></canvas>
        </div>
      </div>

      <div class="results-card">
        <div class="results-head">
          <div class="filter-row">
            <button class="filter-btn active" data-filter="all" onclick="setFilter('all')">All</button>
            <button class="filter-btn" data-filter="success" onclick="setFilter('success')">Success</button>
            <button class="filter-btn" data-filter="redirect" onclick="setFilter('redirect')">Redirects</button>
            <button class="filter-btn" data-filter="client_error" onclick="setFilter('client_error')">4xx</button>
            <button class="filter-btn" data-filter="server_error" onclick="setFilter('server_error')">5xx</button>
            <button class="filter-btn" data-filter="error" onclick="setFilter('error')">Unreachable</button>
            <button class="filter-btn" data-filter="unverifiable" onclick="setFilter('unverifiable')">Unverifiable</button>
          </div>
          <div class="export-row">
            <button class="btn-export" id="exportCsvBtn" onclick="exportCsv()">⬇ Export CSV</button>
            <button class="btn-export" id="exportExcelBtn" onclick="exportExcel()">⬇ Export Excel</button>
          </div>
        </div>
        <div style="overflow-x:auto;">
          <table>
            <thead>
              <tr><th>URL</th><th>Status</th><th>Final URL</th><th>Redirects</th></tr>
            </thead>
            <tbody id="resultsBody"></tbody>
          </table>
        </div>
        <div class="empty-state" id="emptyState" style="display:none;">No URLs match this filter.</div>
      </div>
    </div>
  </main>
</div>

<script>
let currentFilter = 'all';
let allResults = [];
let statusChartInstance = null;
let codesChartInstance = null;

function setFilter(filter) {
  currentFilter = filter;
  document.querySelectorAll('.filter-btn').forEach(b => b.classList.toggle('active', b.dataset.filter === filter));
  renderTable();
}

async function startCheck() {
  const input = document.getElementById('urlInput').value.trim();
  const errorBox = document.getElementById('errorBox');
  errorBox.style.display = 'none';

  if (!input) {
    errorBox.textContent = 'Please paste at least one URL.';
    errorBox.style.display = 'block';
    return;
  }

  const btn = document.getElementById('checkBtn');
  const status = document.getElementById('scanStatus');
  btn.disabled = true;
  btn.innerHTML = '<span class="spinner"></span>Checking…';
  status.textContent = 'This can take a little while for a large batch.';
  document.getElementById('results').style.display = 'none';

  try {
    const resp = await fetch('/api/check', {
      method: 'POST',
      headers: { 'Content-Type': 'application/json' },
      body: JSON.stringify({ input: input })
    });
    const data = await resp.json();

    if (!resp.ok) {
      errorBox.textContent = data.error || 'Something went wrong. Please try again.';
      errorBox.style.display = 'block';
      return;
    }

    allResults = data.results;
    renderSummary(data.summary);
    renderCharts(data.summary);
    currentFilter = 'all';
    document.querySelectorAll('.filter-btn').forEach(b => b.classList.toggle('active', b.dataset.filter === 'all'));
    renderTable();
    document.getElementById('results').style.display = 'block';
    status.textContent = `Checked ${data.summary.total} URL(s).` + (data.summary.truncated ? ` Only the first {{ max_urls }} were checked.` : '');
  } catch (err) {
    errorBox.textContent = 'Network error — please try again in a moment.';
    errorBox.style.display = 'block';
  } finally {
    btn.disabled = false;
    btn.textContent = 'Check Status Codes';
  }
}

function renderSummary(summary) {
  const errors = summary.client_error + summary.server_error;
  document.getElementById('statTotal').textContent = summary.total;
  document.getElementById('statSuccess').textContent = summary.success;
  document.getElementById('statRedirect').textContent = summary.redirect;
  document.getElementById('statErrors').textContent = errors;
  document.getElementById('statUnverifiable').textContent = summary.unverifiable || 0;
  document.getElementById('healthScoreText').textContent = summary.health_score + '%';

  const circumference = 264;
  const offset = circumference - (summary.health_score / 100) * circumference;
  const ring = document.getElementById('healthRing');
  ring.style.strokeDashoffset = offset;
  ring.style.stroke = summary.health_score >= 90 ? '#6af7c8' : summary.health_score >= 70 ? '#f7a26a' : '#f76a6a';
}

function renderCharts(summary) {
  const statusCtx = document.getElementById('statusChart');
  const codesCtx = document.getElementById('codesChart');

  if (statusChartInstance) statusChartInstance.destroy();
  if (codesChartInstance) codesChartInstance.destroy();

  const chartFont = { family: 'Inter', size: 11 };

  statusChartInstance = new Chart(statusCtx, {
    type: 'doughnut',
    data: {
      labels: ['Success', 'Redirects', '4xx', '5xx', 'Unreachable', 'Unverifiable'],
      datasets: [{
        data: [summary.success, summary.redirect, summary.client_error, summary.server_error, summary.error, summary.unverifiable || 0],
        backgroundColor: ['#6af7c8', '#f7a26a', '#f76a6a', '#c94a4a', '#8f3a3a', '#8888a8'],
        borderWidth: 0
      }]
    },
    options: { plugins: { legend: { position: 'bottom', labels: { color: '#c8c8e0', font: chartFont, padding: 10 } } }, cutout: '65%' }
  });

  const codeLabels = (summary.top_status_codes || []).map(c => c[0]);
  const codeValues = (summary.top_status_codes || []).map(c => c[1]);
  codesChartInstance = new Chart(codesCtx, {
    type: 'bar',
    data: {
      labels: codeLabels,
      datasets: [{ data: codeValues, backgroundColor: '#7c6af7', borderRadius: 4 }]
    },
    options: {
      plugins: { legend: { display: false } },
      scales: {
        x: { ticks: { color: '#c8c8e0', font: chartFont }, grid: { display: false } },
        y: { ticks: { color: '#c8c8e0', font: chartFont, precision: 0 }, grid: { color: 'rgba(255,255,255,0.05)' } }
      }
    }
  });
}

function renderTable() {
  const tbody = document.getElementById('resultsBody');
  const emptyState = document.getElementById('emptyState');
  const filtered = currentFilter === 'all' ? allResults : allResults.filter(r => r.severity === currentFilter);

  tbody.innerHTML = '';
  emptyState.style.display = filtered.length ? 'none' : 'block';

  filtered.forEach(r => {
    const tr = document.createElement('tr');
    const statusLabel = (r.error ? r.error : (r.status_code ?? '—')) + (r.redirect_count ? ` (${r.redirect_count} hop${r.redirect_count > 1 ? 's' : ''})` : '');
    const chain = (r.redirect_chain || []).map(h => `${h.status_code} → ${h.url}`).join('<br>');
    const finalCell = r.final_url
      ? `<a href="${r.final_url}" target="_blank" rel="noopener">${r.final_url}</a>${chain ? `<div class="redirect-chain">${chain}</div>` : ''}`
      : '—';

    tr.innerHTML = `
      <td class="url-cell"><a href="${r.url}" target="_blank" rel="noopener">${r.url}</a></td>
      <td><span class="badge ${r.severity}">${statusLabel}</span></td>
      <td class="url-cell">${finalCell}</td>
      <td>${r.redirect_count || 0}</td>
    `;
    tbody.appendChild(tr);
  });
}

function exportCsv() {
  if (!allResults.length) return;
  const rows = [['URL', 'Status Code', 'Category', 'Final URL', 'Redirect Count', 'Redirect Chain', 'Error']];
  allResults.forEach(r => {
    const chain = (r.redirect_chain || []).map(h => `${h.status_code} ${h.url}`).join(' -> ');
    rows.push([
      r.url, r.status_code ?? '', r.severity, r.final_url ?? '', r.redirect_count || 0, chain, r.error ?? ''
    ]);
  });
  const csv = rows.map(row => row.map(cell => `"${String(cell).replace(/"/g, '""')}"`).join(',')).join('\\n');
  const blob = new Blob([csv], { type: 'text/csv' });
  const link = document.createElement('a');
  link.href = URL.createObjectURL(blob);
  link.download = 'http-status-report.csv';
  link.click();
}

async function exportExcel() {
  if (!allResults.length) return;
  const btn = document.getElementById('exportExcelBtn');
  const originalText = btn.textContent;
  btn.disabled = true;
  btn.textContent = 'Preparing…';
  try {
    const resp = await fetch('/api/export/excel', {
      method: 'POST',
      headers: { 'Content-Type': 'application/json' },
      body: JSON.stringify({ results: allResults })
    });
    if (!resp.ok) throw new Error('Export failed');
    const blob = await resp.blob();
    const link = document.createElement('a');
    link.href = URL.createObjectURL(blob);
    link.download = 'http-status-report.xlsx';
    link.click();
  } catch (err) {
    alert('Could not generate the Excel file. Please try again.');
  } finally {
    btn.disabled = false;
    btn.textContent = originalText;
  }
}
</script>
</body>
</html>"""

if __name__ == "__main__":
    app.run(debug=True, port=5051)
